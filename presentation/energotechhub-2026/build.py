"""Editable Russian pitch deck; generation and verification run in Actions only."""
from __future__ import annotations

import hashlib
import json
import os
from pathlib import Path
import re
import sys

if os.environ.get('GITHUB_ACTIONS') != 'true':
    raise SystemExit('Presentation generation is allowed only in GitHub Actions')

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

OUT = Path('artifacts/energotechhub-2026')
OUT.mkdir(parents=True, exist_ok=True)
STEM = 'Sheaft_Energotechhub_2026_RU'
ROOT = 'https://github.com/a-a-k/sheaft-tsfg-experiments'
REPORT = ROOT + '/blob/main/docs/results/revision-v2/final/FINAL_REPORT_RU.md'
PROGRAM = 'https://etechhubspb.ru/accelerator'
BFG = 'https://bfg.ai/bfg-aps/'
AWARD = 'https://conf.researchr.org/info/icse-2026/awards'
FONT = 'DejaVu Sans'
FONT_DIR = Path('/usr/share/fonts/truetype/dejavu')
C = dict(navy='102831', dark='0A1A21', paper='F4F4ED', white='FFFFFF',
         mint='A6E2C6', teal='176B63', orange='E98D52', ink='15343C',
         muted='526C72', dim='AEC5C8', line='C9D7D2', soft='E6ECE5',
         card='193943', pale='DAEEE2')
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(16), Inches(9)
prs.core_properties.title = 'Sheaft — Энерготехнохаб Петербург · Осень 2026'
prs.core_properties.subject = 'Быстрое производственное планирование и проверка исполнения'
prs.core_properties.author = 'Sheaft'
prs.core_properties.keywords = 'Sheaft, APS, ТЭК, Энерготехнохаб, 2026'
manifest = []


def rgb(value):
    return RGBColor.from_string(C.get(value, value))


def rect(s, x, y, w, h, color, radius=False, stroke=None):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                            Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = rgb(color)
    if stroke:
        sh.line.color.rgb = rgb(stroke)
        sh.line.width = Pt(1)
    else:
        sh.line.fill.background()
    if radius:
        sh.adjustments[0] = .1
    return sh


def line(s, x1, y1, x2, y2, color='line', width=1.4):
    sh = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    sh.line.color.rgb = rgb(color)
    sh.line.width = Pt(width)
    return sh


def wrap(text, size, width, bold):
    font = ImageFont.truetype(str(FONT_DIR / ('DejaVuSans-Bold.ttf' if bold else 'DejaVuSans.ttf')), size * 4)
    lines = []
    for paragraph in text.split('\n'):
        words, current = paragraph.split(), ''
        for word in words:
            candidate = (current + ' ' + word).strip()
            if current and font.getlength(candidate) / 4 > width * 72:
                lines.append(current)
                current = word
            else:
                current = candidate
        lines.append(current)
    return lines


def text(s, value, x, y, w, h, size=22, color='ink', bold=False, align=None, link=None):
    assert x >= 0 and y >= 0 and x+w <= 16.01 and y+h <= 9.01, (value, x, y, w, h)
    actual = size
    while actual > 12:
        lines = wrap(value, actual, w-.03, bold)
        if len(lines) * actual * 1.19 <= h*72:
            break
        actual -= 1
    assert len(wrap(value, actual, w-.03, bold))*actual*1.19 <= h*72+2, (value, actual, h)
    box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for i, row in enumerate(wrap(value, actual, w-.03, bold)):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = row
        p.font.name, p.font.size, p.font.bold = FONT, Pt(actual), bold
        p.font.color.rgb = rgb(color)
        p.space_before, p.space_after, p.line_spacing = Pt(0), Pt(0), 1.04
        if align is not None:
            p.alignment = align
    if link:
        # A shape link preserves the authored text colour in LibreOffice PDF.
        # Run-level links are forcibly recoloured blue by that exporter.
        box.click_action.hyperlink.address = link
    manifest[-1]['texts'].append(value)
    return box


def pill(s, value, x, y, w, dark=False, tone='teal'):
    rect(s, x, y, w, .38, 'card' if dark else 'pale', True)
    text(s, value, x+.14, y+.055, w-.28, .26, 12, 'mint' if dark else tone, True)


def base(section, title=None, dark=False, note='', sources=()):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb('navy' if dark else 'paper')
    manifest.append(dict(number=len(prs.slides), section=section, title=title or 'Sheaft',
                         texts=[], sources=list(sources), notes=note))
    text(s, 'SHEAFT / MB3R LAB', .65, .35, 3.45, .34, 15, 'mint' if dark else 'teal', True)
    text(s, section.upper(), 4, .38, 11.3, .3, 12, 'dim' if dark else 'muted', align=PP_ALIGN.RIGHT)
    if title:
        text(s, title, .65, 1.0, 14.65, 1.2, 36, 'white' if dark else 'ink', True)
    line(s, .65, 8.35, 15.35, 8.35, 'card' if dark else 'line', .8)
    text(s, 'Энерготехнохаб Петербург · Осень 2026', .65, 8.55, 11.5, .24, 10, 'dim' if dark else 'muted')
    text(s, f'{len(prs.slides):02}', 14.5, 8.48, .8, .36, 16, 'mint' if dark else 'teal', True, PP_ALIGN.RIGHT)
    if note or sources:
        s.notes_slide.notes_text_frame.text = note + '\n\nИсточники:\n' + '\n'.join(sources)
    return s


def foot(s, value, dark=False, link=None):
    text(s, value, .65, 7.8, 14.7, .42, 11, 'dim' if dark else 'muted', link=link)


def card(s, x, y, w, h, title, body, num=None, dark=False):
    rect(s, x, y, w, h, 'card' if dark else 'white', True)
    offset = .35
    if num:
        text(s, num, x+.3, y+.25, w-.6, .5, 24, 'mint' if dark else 'teal', True)
        offset = .95
    text(s, title, x+.3, y+offset, w-.6, .87, 23, 'white' if dark else 'ink', True)
    text(s, body, x+.3, y+offset+1.05, w-.6, h-offset-1.25, 18, 'dim' if dark else 'muted')


def build():
    s = base('Научно-технологический акселератор', dark=True,
             note='Sheaft помогает составить расписание работ и проверить, что произойдёт при поломке станка или нехватке ресурсов. Предлагаем испытать программу на производственном или ремонтном участке предприятия ТЭК. Программа уже работает на тестовых данных. Показатель 74,73 с получен в модели APS-BASIC; её условия раскрыты на слайде 6. Примерное время: 25 секунд.', sources=[REPORT, PROGRAM])
    pill(s, 'РАБОТАЮЩИЙ ПРОТОТИП', .65, 1.32, 3.5, True)
    text(s, 'Производственный план,\nкоторый можно быстро\nпересчитать и проверить', .65, 2.13, 10.8, 2.9, 43, 'white', True)
    text(s, 'Sheaft составляет расписания и проверяет последствия сбоев на производственных и ремонтных участках ТЭК', .7, 5.62, 10.15, 1.15, 24, 'dim')
    for i, (x, w) in enumerate(((11.8, 2.2), (12.4, 2.8), (11.2, 2.5), (12, 3.0))):
        rect(s, x, 2.2+i*.77, w, .42, 'mint' if i != 2 else 'orange', True)
        line(s, 10.95, 2.7+i*.77, 15.2, 2.7+i*.77, 'card', 1)
    foot(s, 'Предлагаем проверить Sheaft на данных предприятия. Результаты экспериментов опубликованы.', True, REPORT)

    s = base('01 / Задача предприятия', 'Остановка одного станка задерживает следующие работы',
             note='Пример задачи для будущего пилота: станок остановился, детали ждут обработки, следующие операции начинаются позже. Если нескольким работам нужна одна бригада или один стенд, их занятость тоже влияет на срок. Диспетчеру нужно понять, какие заказы задержатся и как изменить расписание. Примерное время: 30 секунд.')
    text(s, 'Нужно понять, какие заказы задержатся и как изменить расписание.', .7, 2.36, 14.5, .6, 24)
    for x, title, label, tone in [(1.0, 'Операция 1', 'обработка', 'teal'), (6.0, 'Операция 2', 'ожидание', 'orange'), (11, 'Операция 3', 'сдвиг срока', 'teal')]:
        rect(s, x, 3.67, 4, 1.55, 'white', True)
        rect(s, x, 3.67, .09, 1.55, tone)
        text(s, title, x+.27, 3.98, 3.5, .45, 23, bold=True)
        text(s, label, x+.27, 4.57, 3.5, .32, 16, tone)
    for x in (5.2, 10.2):
        text(s, '→', x, 4.1, .65, .6, 28, 'teal')
    rect(s, .7, 6.05, 14.6, 1.17, 'navy', True)
    text(s, 'Диспетчер сравнивает варианты и выбирает новый порядок работ.', 1.03, 6.33, 13.9, .65, 25, 'white', True)
    foot(s, 'Пример задачи для пилота. Влияние на сроки и затраты проверим на данных предприятия.')

    s = base('02 / Где начнём', 'Предлагаем начать с участка изготовления или ремонта деталей',
             note='Для первого испытания подойдёт изготовление деталей или восстановление узлов энергетического оборудования. На таком участке можно описать порядок операций, доступные станки, бригады и места для ожидания деталей. Участок и данные выбираем вместе с предприятием. Примерное время: 30 секунд.', sources=[PROGRAM])
    card(s, .7, 2.67, 4.6, 4.3, 'Пользователи', 'Диспетчер и плановый отдел\n\nРезультат оценивает руководитель участка', '01')
    card(s, 5.7, 2.67, 4.6, 4.3, 'Работы', 'Изготовление деталей и ремонт узлов\n\nУчитываем станки, бригады и места для деталей', '02')
    card(s, 10.7, 2.67, 4.6, 4.3, 'Результат', 'Новый порядок работ после сбоя\n\nРасчёт сроков и причин задержки', '03')
    foot(s, 'Первый пилот проведём на выбранном участке и сравним расчёт с фактическим ходом работ.')

    s = base('03 / Как работает Sheaft', 'Составляем расписание, проверяем сбои, сравниваем варианты', dark=True,
             note='Так будет устроена работа диспетчера с продуктом. Программы уже составляют расписания и рассчитывают их выполнение при заданных сбоях. Интерфейс и обмен данными с системами предприятия предстоит разработать. Сейчас расчёт выявляет взаимную блокировку операций; автоматическое изменение очередей для её устранения требует доработки. Примерное время: 35 секунд.', sources=[REPORT])
    steps = [('Загрузить', 'Заказы, порядок работ\nи доступные станки'), ('Составить', 'На каком станке\nи когда выполнить работу'), ('Проверить', 'Какие заказы задержит\nполомка или нехватка мест'), ('Выбрать', 'Сравнить варианты\nи утвердить расписание')]
    for i, (title, body) in enumerate(steps):
        x = .7+i*3.75
        card(s, x, 2.83, 3.4, 3.6, title, body, str(i+1).zfill(2), True)
        if i < 3:
            text(s, '→', x+3.42, 4.07, .35, .5, 21, 'mint')
    text(s, 'Примеры изменений: поломка станка · срочный заказ · занятая бригада', .75, 6.93, 14.5, .55, 21, 'mint')
    foot(s, 'Расчётные программы работают. Интерфейс и обмен данными с системами предприятия предстоит разработать.', True)

    s = base('04 / Исследовательская основа', 'Модель показывает, как отказ влияет на выполнение задачи',
             note='Исследование Sheaft начиналось с программных систем: восстанавливаем связи по данным об их работе, затем рассчитываем последствия отказов. Статья Model Discovery and Graph Simulation: A Lightweight Gateway to Chaos Engineering получила NIER 2026 Distinguished Paper Award; автор Anatoly Krasnovsky. Для производства отдельно проверяем порядок операций, ожидание деталей и занятость ресурсов. Стандартные алгоритмы LIST, DAG и DES сами по себе не являются нашей научной новизной. Примерное время: 35 секунд.', sources=[AWARD,REPORT])
    rect(s,.7,2.7,7.05,4.45,'navy',True)
    pill(s,'ИССЛЕДОВАНО В ИТ',1.03,3.03,3.9,True)
    text(s,'Строим модель по данным\nи рассчитываем последствия отказов',1.06,3.82,6.33,1.2,29,'white',True)
    text(s,'ICSE 2026 · NIER',1.06,5.42,6.3,.57,27,'mint',True)
    text(s,'Distinguished Paper Award',1.06,6.15,6.3,.56,21,'dim')
    text(s,'ПРИМЕНЕНИЕ НА ПРОИЗВОДСТВЕ',8.33,2.92,6.7,.4,17,'teal',True)
    text(s,'Проверяем, как остановка станка меняет сроки следующих операций.',8.33,3.8,6.7,1.35,29,bold=True)
    text(s,'Расчёт проверен на тестовых задачах. Далее сопоставим его с работой реального участка.',8.33,5.65,6.7,1.13,22,'muted')
    foot(s,'Награда ICSE относится к исследованию программных систем. Производство проверяем в отдельной серии опытов.',link=AWARD)

    s = base('05 / Измеренный результат', 'Миллион операций — за 74,73 секунды', dark=True,
             note='Это максимум медиан по десяти уникальным миллионным входам, каждый с тремя повторами. На двух размерах выполнено 60 основных запусков; все планы прошли проверку допустимости. Время включает запуск, чтение, построение и полный вывод, но не последующую валидацию. Ограничения измеренной модели: альтернативные станки, предшествования, поступления; буферы неограничены, общие ресурсы не учитываются в построении. Оптимальность не заявлена. Примерное время: 40 секунд.', sources=[REPORT, ROOT+'/blob/main/docs/results/revision-v2/final/APS_MATRIX_REPORT_RU.md'])
    text(s, '74,73', .7, 2.58, 9.5, 1.95, 110, 'mint', True)
    text(s, 'секунды на построение и полный вывод плана', .9, 4.88, 8.55, 1.18, 26, 'white')
    for y, n, label in [(2.65, '60', 'запусков: все планы прошли проверку'), (4.05, '10 × 3', 'задач по миллиону операций × повторы'), (5.55, '1 CPU', 'выделено для расчёта')]:
        text(s, n, 11.0, y, 4.15, .7, 37, 'white', True)
        text(s, label, 11.03, y+.75, 4.05, .62, 17, 'dim')
    foot(s, 'Тестовые задачи APS-BASIC. Места для деталей не ограничены; конкуренция за бригады и оснастку не учитывалась.', True, REPORT)

    s = base('06 / Причины ожидания', 'Учитываем места для деталей и занятость бригад',
             note='В модели одновременно учитываются два условия. Первое: количество мест, где деталь ждёт следующей операции, ограничено. Если передать готовую деталь некуда, она продолжает занимать станок. Второе: нескольким операциям может требоваться одна бригада или единица оснастки; одновременно занять её нельзя. Эта проверка прошла 135 сценариев Mk01 и 405 проверок состояния к сроку. В тестах есть простой последовательный план. Далее нужно научить планировщик менять очереди с учётом обоих условий и проверить, насколько хороший план получается. Примерное время: 40 секунд.', sources=[REPORT, ROOT+'/blob/main/docs/PBR_EXACT_V2_2_RU.md'])
    text(s,'Места между станками заняты',.8,2.68,7.0,.83,26,bold=True)
    text(s,'Готовая деталь удерживает станок, пока её некуда передать.',.8,3.55,6.8,.95,23,'muted')
    text(s,'Одна бригада нужна двум работам',.8,4.92,7.0,.83,26,bold=True)
    text(s,'Следующая работа ждёт, пока бригада закончит предыдущую.',.8,5.81,6.8,.95,23,'muted')
    rect(s,8.25,2.6,7.05,4.6,'navy',True)
    text(s,'Проверка уже работает',8.65,3.02,6.2,.76,26,'mint',True)
    text(s,'135 сценариев\n405 проверок к сроку',8.65,4.04,6.15,1.15,30,'white',True)
    text(s,'Далее — автоматически менять очереди, чтобы сократить такое ожидание.',8.65,5.7,6.15,1.05,23,'dim')
    foot(s,'Результаты проверки на задаче Mk01. Скорость и качество нового построителя расписаний ещё предстоит измерить.',link=REPORT)

    s = base('07 / Сравнение с APS', 'Проверим, быстрее ли Sheaft составляет производственный план',
             note='Сохраняем согласованную гипотезу преимущества перед классом APS с BFG как первым ориентиром. Поставщик публикует миллион операций за 45 минут. Наши 74,73 секунды получены в другой модели и среде. Сравнение качества, ограничений и аппаратуры ещё нужно выровнять. Число 36× не используем как доказанное конкурентное преимущество. Примерное время: 30 секунд.', sources=[BFG, REPORT])
    for x, title, value, detail in [(.7, 'BFG APS · данные поставщика', '45 минут', 'На сайте указано время построения\nплана на миллион операций'), (8.2, 'Sheaft · наш эксперимент', '74,73 с', 'Миллион операций в модели APS-BASIC\nна тестовых задачах')]:
        rect(s, x, 2.7, 7.05, 3.25, 'white', True)
        text(s, title, x+.33, 3.02, 6.35, .62, 23, bold=True)
        text(s, value, x+.33, 3.93, 6.35, .9, 45, 'teal', True)
        text(s, detail, x+.33, 5.04, 6.35, .73, 18, 'muted')
    text(s, 'Следующий опыт: одинаковые заказы, требования к плану и оборудование для расчёта.', .78, 6.38, 14.45, .96, 26, bold=True)
    foot(s, 'Условия двух измерений различаются. Преимущество по скорости при одинаковом качестве пока не доказано.', link=BFG)

    s = base('08 / Как оценим пилот', 'Сравним Sheaft с текущим способом планирования на участке',
             note='Нужен участок, по которому сохранились расписания и данные о выполненных работах. До расчётов согласуем правила модели и критерии сравнения с текущим способом планирования. После этого проверим новые расписания на одних и тех же заказах и сбоях. Сравним выпуск к сроку, задержки, время расчёта и объём ручной работы. Примерное время: 40 секунд.')
    rows = [('Выполнимость', 'Нет конфликтов за станки, бригады и места для деталей.'), ('Сроки заказов', 'Сравниваем готовый объём к сроку и длительность задержек.'), ('Время расчёта', 'Диспетчер получает результат за согласованное время.'), ('Работа диспетчера', 'Считаем ручные исправления и время на подготовку расписания.')]
    for i, (title, desc) in enumerate(rows):
        y = 2.63+i*1.12
        text(s, str(i+1).zfill(2), .75, y+.05, .8, .48, 23, 'teal', True)
        text(s, title, 1.85, y, 3.85, .72, 23, bold=True)
        text(s, desc, 6.15, y+.02, 8.72, .83, 21, 'muted')
        if i < 3:
            line(s, 1.85, y+.94, 15, y+.94)
    foot(s, 'До начала пилота согласуем данные для проверки и значения, при которых результат считается успешным.')

    s = base('09 / План внедрения', 'Сначала проверим расчёт, затем подключим его к работе участка', dark=True,
             note='Начинаем с выгрузки данных и описания участка. Дорабатываем планировщик, чтобы он учитывал фактические правила работы. Затем диспетчер сравнивает предлагаемые расписания с текущими, сохраняя решение за собой. После приёмки подключаем обмен данными с системами предприятия. Сроки определим после просмотра данных и согласования объёма работ. Предлагаем оплачиваемый пилот или НИОКР, затем лицензию и сопровождение. Примерное время: 35 секунд.', sources=[PROGRAM])
    line(s, 1.3, 3.03, 14.55, 3.03, 'mint', 2)
    stages = [('Описать участок', 'Порядок работ, станки,\nбригады и места для деталей'), ('Доработать расчёт', 'Учесть правила работы\nвыбранного участка'), ('Проверить в работе', 'Диспетчер сравнивает\nрасписания и выбирает вариант'), ('Подключить данные', 'Обмен с системами учёта\nи управления производством')]
    for i, (title, desc) in enumerate(stages):
        x = .75+i*3.77
        dot = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x+.12), Inches(2.81), Inches(.44), Inches(.44))
        dot.fill.solid(); dot.fill.fore_color.rgb=rgb('orange' if i==1 else 'mint'); dot.line.fill.background()
        text(s, str(i+1).zfill(2), x, 3.7, 3.3, .62, 28, 'mint', True)
        text(s, title, x, 4.53, 3.32, .97, 25, 'white', True)
        text(s, desc, x, 5.72, 3.25, 1.28, 18, 'dim')
    foot(s, 'Предлагаем оплачиваемый пилот или НИОКР, затем — лицензию на программу и сопровождение.', True)

    s = base('10 / Команда MB3R Lab', 'Исследования, разработка и работа с заказчиком',
             note='Состав и роли перенесены по указанию пользователя из презентации Sheaft для демо-дня «Большой разведки» 2026. Анатолий Красновский — основатель и исследователь, полная занятость. Yehia Sobeh — разработка и безопасность, частичная занятость. Руслан Саяхов — советник по продукту, 10+ лет в B2B-продуктах. Награда Анатолия подтверждена официальным списком ICSE. Для пилота нужен технолог и владелец данных со стороны предприятия. Примерное время: 25 секунд.',sources=[AWARD])
    card(s,.7,2.65,4.58,4.25,'Анатолий\nКрасновский','Основатель и исследователь\n\n10+ лет в разработке ПО\nМетод и реализация Sheaft\nICSE 2026 NIER — награда','ПОЛНАЯ ЗАНЯТОСТЬ')
    card(s,5.71,2.65,4.58,4.25,'Yehia\nSobeh','Разработка и безопасность\n\nСерверная реализация\nПодготовка поставки\nДанные и адаптеры пилота','ЧАСТИЧНАЯ ЗАНЯТОСТЬ')
    card(s,10.72,2.65,4.58,4.25,'Руслан\nСаяхов','Советник по продукту\n\n10+ лет в B2B-продуктах\nСбер, ЦИАН, Ozon.Tech\nПилот и переход к договору','ПРОДУКТОВАЯ ЭКСПЕРТИЗА')
    foot(s,'Для пилота нужен технолог предприятия и сотрудник, который поможет подготовить данные.')

    s = base('11 / Запрос к акселератору', dark=True,
             note='Ищем предприятие, готовое проверить Sheaft на одном производственном или ремонтном участке. Нужны данные о заказах и выполненных работах, а также технолог или диспетчер, который поможет проверить модель. После просмотра данных согласуем задачи, сроки и стоимость пилота. Примерное время: 25 секунд.', sources=[ROOT, PROGRAM])
    text(s, 'Ищем предприятие,\nчтобы проверить Sheaft\nна реальных заказах', .7, 1.68, 11.4, 2.5, 46, 'white', True)
    for i, label in enumerate(('Участок изготовления или ремонта деталей', 'Данные о заказах и выполненных работах', 'Технолог или диспетчер для проверки результата')):
        text(s, '0'+str(i+1), .8, 4.77+i*.78, .8, .47, 21, 'mint', True)
        text(s, label, 1.85, 4.77+i*.78, 12.6, .56, 24, 'white')
    text(s,'Анатолий Красновский · основатель MB3R Lab',.8,7.25,9.55,.4,18,'dim')
    text(s,'contact@mb3r-lab.org',10.52,7.25,4.78,.4,18,'mint',link='mailto:contact@mb3r-lab.org')
    foot(s, 'Открытые результаты: github.com/a-a-k/sheaft-tsfg-experiments', True, ROOT)

    s = base('Приложение A / Алгоритмы', 'LIST строит расписание, DAG и DES проверяют его выполнение',
             note='LIST выбирает назначения и очереди; DAG рассчитывает времена готового плана при неограниченных буферах; DES проверяет исполнение с конечными буферами и общими ресурсами. Независимость переходов позволяет сопоставлять результаты. TSFG-агрегация остаётся отдельным исследовательским направлением.', sources=[REPORT])
    card(s,.7,2.8,4.55,3.9,'LIST','Строит новый план\n\nВыбирает оборудование и порядок операций','ПЛАНИРОВАНИЕ')
    card(s,5.72,2.8,4.55,3.9,'DAG','Проверяет готовый план\n\nФиксированные очереди, неограниченные буферы','БЫСТРЫЙ РАСЧЁТ')
    card(s,10.74,2.8,4.55,3.9,'DES','Моделирует исполнение\n\nКонечные буферы и общие ресурсы','ОГРАНИЧЕНИЯ')
    text(s,'Сравниваем результаты независимых программ. Методика и данные опубликованы.',.75,7.03,14.3,.48,23,'teal',True)
    foot(s,'Потоковая агрегация TSFG — отдельное направление для совместимых непрерывных процессов.')

    s = base('Приложение B / Готовность', 'Что уже работает и что предстоит проверить',
             note='Эту таблицу использовать в ответах экспертов. Состояния относятся к текущему прототипу, а не к готовой промышленной поставке. Формальный TRL без согласованной оценки не назначаем.', sources=[REPORT])
    statuses = [('Построение расписания APS-BASIC', 'Работает; скорость измерена', '60 запусков; планы прошли проверку'), ('Проверка мест ожидания и ресурсов', 'Работает; результаты сверены', '135 сценариев; независимые программы'), ('Выполнение заказов по одному', 'Простой вариант для тестов', 'Может увеличить срок всего задания'), ('Перестройка плана при нехватке ресурсов', 'Нужно доработать', 'Скорость и качество предстоит измерить'), ('Интерфейс и обмен данными', 'Нужно разработать', 'Для работы на предприятии')]
    for i, (a,b,c) in enumerate(statuses):
        y=2.53+i*.98
        if i%2==0:rect(s,.7,y-.08,14.6,.89,'white',True)
        text(s,a,.94,y+.07,5.1,.66,18,bold=True)
        text(s,b,6.38,y+.07,3.7,.66,17,'teal',True)
        text(s,c,10.47,y+.07,4.48,.66,16,'muted')
    foot(s, 'Часть опытов не завершена. Причины остановок и все полученные результаты сохранены в отчёте.', link=REPORT)

    s = base('Приложение C / Методика', 'Что означает показатель 74,73 секунды',
             note='Для каждого размера — 100 тысяч и миллион операций — взяли 10 разных задач и выполнили каждую трижды. Для каждой задачи выбрали средний по порядку результат трёх запусков. Самый большой из этих результатов для миллиона операций — 74,73 с. Время включает чтение данных, построение и запись расписания. Проверка допустимости выполнялась после измерения. Опыт не измеряет скорость построения расписания с ограниченным количеством мест для деталей и общими бригадами.', sources=[REPORT, ROOT+'/blob/main/docs/results/revision-v2/final/APS_MATRIX_REPORT_RU.md'])
    card(s,.7,2.67,7.05,3.95,'Как получили результат','10 разных задач × 3 повтора на каждом размере: 100 тысяч и миллион операций.\n\n74,73 с — максимум медиан для миллиона.','ЗАДАЧИ И ПОВТОРЫ')
    card(s,8.2,2.67,7.05,3.95,'Что входит во время','Запуск процесса, чтение входа, назначения, очереди и полный вывод.\n\nPython 3.12 · 1 CPU · 4 GiB · GitHub Actions.','УСЛОВИЯ')
    text(s,'Все планы прошли проверку. Опыт не доказывает, что найдено лучшее расписание.',.8,7.02,14.4,.54,24,'teal',True)
    foot(s,'APS-BASIC учитывает поступления заказов, порядок операций и выбор станка. Условия полностью описаны в отчёте.',link=REPORT)

    s = base('Приложение D / Остальные результаты', 'Какие предположения эксперимент пока не подтвердил',
             note='Не скрывать отрицательные результаты. В F1/100k K=100 завершились 18 DES-процессов, все 18 TSFG остановлены по 300 с. Общие завершённые траектории совпали, полных пар нет; гипотеза ускорения INCONCLUSIVE. Большие входы были слабыми контролями совместной связи. Два банка рекомендаций по 1000 сценариев не прошли критерий полезности; шесть других областей не завершены/не допущены. Потоковая агрегация отдельна от точных операций.', sources=[REPORT])
    card(s,.7,2.67,4.58,4.38,'Точное исполнение','Преимущество TSFG по скорости не установлено.\n\nВ F1/100k: 18 DES завершены; 18 TSFG — таймаут 300 с.','01')
    card(s,5.71,2.67,4.58,4.38,'Выбор усиления','Выбранное усиление станка не прошло проверку пользы.\n\nДве завершённые выборки по 1 000 сценариев.','02')
    card(s,10.72,2.67,4.58,4.38,'Потоковая модель','Шесть примеров прошли проверку точности.\n\nСкорость расчёта и частота ошибок при случайных сбоях не измерены.','03')
    foot(s,'Отчёт содержит исходные данные, результаты и причины незавершённых опытов.',link=REPORT)

    s = base('Приложение E / Источники', 'Публикации, данные и описание программы',
             note='Все ссылки кликабельны в PPTX и PDF. Материалы программы, поставщика и ICSE проверены 27 сентября 2026. Команда и контакт взяты из прежних презентаций Sheaft по указанию пользователя: «Промышленный пилот» и «Большая разведка — демо-день 2026». Исходные презентации и личные документы не опубликованы в экспериментальном репозитории.', sources=[PROGRAM,REPORT,BFG,AWARD,ROOT+'/releases/tag/revision-v2-final-package-16366e27dea4'])
    sources = [('01','Программа акселератора','Помощь в подготовке и внедрении технологий на предприятиях ТЭК.',PROGRAM),('02','Экспериментальный отчёт Sheaft v2.2','Как проводили опыты, что получили и что осталось непроверенным.',REPORT),('03','Открытый архив данных','Исходные задачи, версии программ и результаты расчётов.',ROOT+'/releases/tag/revision-v2-final-package-16366e27dea4'),('04','BFG APS — данные поставщика','Миллион операций за 45 минут; условия отличаются от нашего опыта.',BFG)]
    for i,(n,title,desc,url) in enumerate(sources):
        y=2.55+i*1.18
        text(s,n,.8,y,.75,.4,22,'teal',True)
        text(s,title,1.87,y,12.85,.49,24,'teal',True,link=url)
        text(s,desc,1.87,y+.6,12.85,.43,18,'muted')
    text(s,'05',.8,7.23,.75,.4,22,'teal',True)
    text(s,'ICSE 2026 NIER — официальный список наград',1.87,7.23,12.85,.45,21,'teal',True,link=AWARD)
    foot(s,'Подготовлено 27 сентября 2026 · 12 основных слайдов + 5 слайдов для вопросов экспертов.')

    prs.save(OUT / f'{STEM}.pptx')
    (OUT/'deck-manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
    notes = ['# Sheaft — заметки докладчика\n\nОсновная часть: 12 слайдов, ориентир 5–7 минут. Приложение — для вопросов.\n']
    for item in manifest:
        notes.append(f"## {item['number']:02}. {item['title']}\n\n{item['notes']}\n\n" + '\n'.join(f'- {u}' for u in item['sources']))
    (OUT/'Заметки_докладчика_RU.md').write_text('\n\n'.join(notes),encoding='utf-8')
    print(json.dumps(dict(status='BUILT',slides=len(prs.slides)),ensure_ascii=False))


def normalized(value):
    return re.sub(r'\s+', '', value).replace('\u00ad','').replace('−','-')


def verify():
    import fitz
    data=json.loads((OUT/'deck-manifest.json').read_text(encoding='utf-8'))
    doc=fitz.open(OUT/f'{STEM}.pdf')
    assert len(doc)==len(data)==17
    pdf_links=[link.get('uri','') for page in doc for link in page.get_links()]
    assert 'mailto:contact@mb3r-lab.org' in pdf_links, pdf_links
    assert AWARD in pdf_links and REPORT in pdf_links
    issues=[]
    thumbs=[]
    previews=OUT/'previews';previews.mkdir(exist_ok=True)
    for i,page in enumerate(doc):
        content=normalized(page.get_text())
        for expected in data[i]['texts']:
            if normalized(expected) not in content:
                issues.append(dict(slide=i+1,kind='text_missing_or_interrupted',text=expected))
        for b in page.get_text('dict')['blocks']:
            for l in b.get('lines',[]):
                for span in l.get('spans',[]):
                    x0,y0,x1,y1=span['bbox']
                    if min(x0,y0)<-.5 or x1>page.rect.width+.5 or y1>page.rect.height+.5:
                        issues.append(dict(slide=i+1,kind='text_outside_page',text=span['text']))
        pix=page.get_pixmap(matrix=fitz.Matrix(1,1),alpha=False)
        pix.save(str(previews/f'slide-{i+1:02}.png'))
        im=Image.frombytes('RGB',(pix.width,pix.height),pix.samples)
        im.thumbnail((480,270))
        thumbs.append(im)
    sheet=Image.new('RGB',(4*500,((len(thumbs)+3)//4)*304),(224,230,226))
    draw=ImageDraw.Draw(sheet)
    for i,im in enumerate(thumbs):
        x=(i%4)*500+10;y=(i//4)*304+10
        sheet.paste(im,(x,y))
        draw.text((x,y+275),f'{i+1:02}',fill=(16,40,49))
    sheet.save(OUT/'Обзор_слайдов.png')
    result=dict(status='PASS' if not issues else 'FAIL',slides=len(doc),issues=issues,links=len(pdf_links),
                source_commit=os.environ.get('GITHUB_SHA'),run_id=os.environ.get('GITHUB_RUN_ID'))
    (OUT/'verification.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    hashes={p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in OUT.iterdir() if p.is_file() and p.name!='SHA256.json'}
    (OUT/'SHA256.json').write_text(json.dumps(hashes,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(result,ensure_ascii=False))
    if issues:raise SystemExit(1)


if __name__=='__main__':
    verify() if '--verify' in sys.argv else build()
